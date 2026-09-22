"""Tests for rtc.exporter."""

from __future__ import annotations

import csv

import openpyxl
import pytest

from rtc.exceptions import UnsupportedFileFormatError
from rtc.exporter import COLUMNS, export_test_cases
from rtc.generator import TestCaseGenerator


@pytest.fixture
def sample_test_cases(simple_requirement):
    generator = TestCaseGenerator()
    return generator.generate_for_requirement(simple_requirement)


def test_export_to_xlsx_creates_readable_workbook(tmp_path, sample_test_cases):
    output_path = tmp_path / "test_cases.xlsx"
    result_path = export_test_cases(sample_test_cases, output_path)

    assert result_path.exists()

    workbook = openpyxl.load_workbook(result_path)
    sheet = workbook.active

    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header == COLUMNS

    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == len(sample_test_cases)


def test_export_to_csv_creates_readable_file(tmp_path, sample_test_cases):
    output_path = tmp_path / "test_cases.csv"
    result_path = export_test_cases(sample_test_cases, output_path)

    assert result_path.exists()

    with result_path.open(newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)

    assert rows[0] == COLUMNS
    assert len(rows) - 1 == len(sample_test_cases)


def test_export_unsupported_extension_raises(tmp_path, sample_test_cases):
    output_path = tmp_path / "test_cases.txt"
    with pytest.raises(UnsupportedFileFormatError):
        export_test_cases(sample_test_cases, output_path)


def test_export_creates_missing_parent_directories(tmp_path, sample_test_cases):
    output_path = tmp_path / "nested" / "dirs" / "test_cases.csv"
    result_path = export_test_cases(sample_test_cases, output_path)
    assert result_path.exists()


def test_export_empty_test_case_list_produces_header_only(tmp_path):
    output_path = tmp_path / "empty.csv"
    result_path = export_test_cases([], output_path)

    with result_path.open(newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows == [COLUMNS]
